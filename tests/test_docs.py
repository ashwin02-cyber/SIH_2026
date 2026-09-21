"""Submission deliverables: video script timing, sample reports, slide deck, docs cross-references."""
import io
import json
import os
import re

import pytest
from pptx import Presentation
from pypdf import PdfReader

from conftest import ROOT

DOCS = os.path.join(ROOT, "docs")


def _video_rows():
    rows = []
    with open(os.path.join(DOCS, "VIDEO_SCRIPT.md"), encoding="utf-8") as f:
        for line in f:
            m = re.match(r"\|\s*(\d+)\s*\|\s*(\d+):(\d+)-(\d+):(\d+)\s*\|(.*?)\|(.*?)\|(.*?)\|", line)
            if m:
                start, end = int(m[2]) * 60 + int(m[3]), int(m[4]) * 60 + int(m[5])
                rows.append({"start": start, "end": end, "narration": m[7].strip()})
    return rows


def test_video_script_is_exactly_three_minutes_and_contiguous():
    rows = _video_rows()
    assert len(rows) >= 6
    assert rows[0]["start"] == 0 and rows[-1]["end"] == 180
    for prev, cur in zip(rows, rows[1:]):
        assert cur["start"] == prev["end"], "shots must be back to back"
    assert all(r["end"] > r["start"] for r in rows)


def test_video_narration_can_actually_be_spoken_in_time():
    rows = _video_rows()
    for r in rows:
        words = len(r["narration"].split())
        assert words / (r["end"] - r["start"]) <= 2.7, f"shot at {r['start']}s has too many words"
    total = sum(len(r["narration"].split()) for r in rows)
    assert 200 <= total <= 480


def test_sample_reports_are_real_and_from_a_real_capture():
    d = os.path.join(DOCS, "sample_reports")
    analysis = json.load(open(os.path.join(d, "analysis.json"), encoding="utf-8"))
    assert analysis["filename"].endswith(".pcap")
    assert os.path.exists(os.path.join(ROOT, "data", "samples", analysis["filename"]))
    for kind in ("executive", "technical"):
        blob = open(os.path.join(d, f"sample-{kind}.pdf"), "rb").read()
        assert blob[:5] == b"%PDF-"
        text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(blob)).pages)
        assert analysis["filename"] in text and analysis["risk_level"] in text
        assert "declared" in text.lower()  # the report discloses that values come from the file name


def test_deck_is_intact():
    prs = Presentation(os.path.join(DOCS, "IPsec_VPN_Analyzer_Deck.pptx"))
    assert len(prs.slides) >= 10
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, f"slide {i}: shape off the slide"
            assert shape.left + shape.width <= prs.slide_width + 10, f"slide {i}: shape past right edge"
            assert shape.top + shape.height <= prs.slide_height + 10, f"slide {i}: shape past bottom edge"
        assert slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip(), f"slide {i}: no speaker notes"
    text = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    assert "TODO" not in text and "lorem" not in text.lower()


def test_deck_numbers_match_metrics_json():
    metrics = json.load(open(os.path.join(ROOT, "ml-engineer", "metrics.json"), encoding="utf-8"))
    prs = Presentation(os.path.join(DOCS, "IPsec_VPN_Analyzer_Deck.pptx"))
    cells = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_table:
                cells += [c.text for r in sh.table.rows for c in r.cells]
    baseline = metrics["evaluation"]["sanity_checks"]["majority_class_window_accuracy"]
    assert f"{baseline * 100:.1f} %" in cells


@pytest.mark.parametrize("path", ["README.md", "DEPLOY.md", "docs/TECHNICAL_DOCUMENTATION.md",
                                  "integration-docs-lead/README.md", "integration-docs-lead/integration-guide.md"])
def test_relative_links_in_docs_resolve(path):
    full = os.path.join(ROOT, path)
    text = open(full, encoding="utf-8").read()
    base = os.path.dirname(full)
    for target in re.findall(r"\]\(([^)#\s]+)\)", text):
        if target.startswith(("http://", "https://", "mailto:")):
            continue
        assert os.path.exists(os.path.normpath(os.path.join(base, target))), f"{path}: broken link {target}"
